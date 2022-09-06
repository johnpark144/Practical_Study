from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# B2:C11까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # 차트 종류 설정(Bar, Line, Pie, ..)
# bar_chart.add_data(bar_value) # 차트 데이터 추가
# ws.add_chart(bar_chart, "E1") # 차트 넣을 위치 정의

# B1:C11 까지의 데이터
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # 계열 > 영어, 수학(제목에서 가져옴)
line_chart.title = "성적표" # 제목
line_chart.style =  10 # 미리 정의된 스타일 적용, 사용자가 개별 지정도 가능
line_chart.y_axis.title = "점수" # Y축의 제목
line_chart.x_axis.title = "번호" # X축의 제목
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")